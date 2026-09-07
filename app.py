"""Bus attendance - Flask backend.

Stores students (with face descriptors computed in the browser) and daily
attendance records in SQLite, and exports attendance as Excel or PDF.
"""
import io
import json
import os
import re
import sqlite3
from datetime import date, datetime

from flask import Flask, g, jsonify, request, send_file, send_from_directory

import arabic_reshaper
from bidi.algorithm import get_display

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.environ.get("DATA_DIR", os.path.join(BASE_DIR, "data"))
os.makedirs(DATA_DIR, exist_ok=True)
DB_PATH = os.path.join(DATA_DIR, "attendance.db")

app = Flask(__name__, static_folder="static", static_url_path="/static")
app.config["MAX_CONTENT_LENGTH"] = 8 * 1024 * 1024  # photos are small thumbnails

DAY_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
FONT_DIR = os.path.join(BASE_DIR, "static", "fonts")

# ---------------------------------------------------------------------- i18n
STRINGS = {
    "ar": {
        "name_required": "الاسم مطلوب",
        "id_required": "رقم الهوية الوطنية مطلوب",
        "photo_image": "يجب أن تكون الصورة ملف صورة",
        "photo_large": "الصورة كبيرة جدًا",
        "id_exists": "رقم الهوية '{id}' مسجّل مسبقًا",
        "not_found": "الطالب غير موجود",
        "record_not_found": "السجل غير موجود",
        "invalid_date": "تاريخ غير صالح",
        "invalid_range": "نطاق تاريخ غير صالح",
        "desc_list": "يجب أن تكون بصمات الوجه قائمة",
        "desc_len": "كل بصمة وجه يجب أن تحتوي على 128 رقمًا",
        "student_id_required": "معرّف الطالب مطلوب",
        "report_title": "تقرير حضور الحافلة",
        "summary_title": "ملخص حضور الحافلة",
        "bus": "الحافلة",
        "period": "الفترة",
        "school_days": "أيام الدراسة المسجلة",
        "students": "الطلاب",
        "generated": "تاريخ الإنشاء",
        "sheet_summary": "الملخص",
        "sheet_daily": "اليومي",
        "sheet_log": "السجل",
        "summary_heading": "الملخص (سجل المسح الكامل موجود في ملف إكسل)",
        "daily_heading": "الحضور اليومي (ح = حاضر، فارغ = غائب)",
        "no": "الرقم",
        "national_id": "رقم الهوية الوطنية",
        "name": "الاسم",
        "grade": "الصف",
        "parent_phone": "هاتف ولي الأمر",
        "days_present": "أيام الحضور",
        "days_absent": "أيام الغياب",
        "pct": "نسبة الحضور %",
        "present": "حاضر",
        "absent": "غائب",
        "total": "الإجمالي",
        "date": "التاريخ",
        "time": "الوقت",
        "method": "الطريقة",
        "method_face": "وجه",
        "method_manual": "يدوي",
        "mark_present": "ح",
        "mark_absent": "غ",
    },
    "en": {
        "name_required": "Name is required",
        "id_required": "National ID is required",
        "photo_image": "Photo must be an image",
        "photo_large": "Photo is too large",
        "id_exists": "National ID '{id}' already exists",
        "not_found": "Student not found",
        "record_not_found": "Record not found",
        "invalid_date": "Invalid date",
        "invalid_range": "Invalid date range",
        "desc_list": "descriptors must be a list",
        "desc_len": "each descriptor must have 128 numbers",
        "student_id_required": "student_id is required",
        "report_title": "Bus Attendance Report",
        "summary_title": "Bus Attendance Summary",
        "bus": "Bus",
        "period": "Period",
        "school_days": "School days with records",
        "students": "Students",
        "generated": "Generated",
        "sheet_summary": "Summary",
        "sheet_daily": "Daily",
        "sheet_log": "Log",
        "summary_heading": "Summary (the full scan log is included in the Excel export)",
        "daily_heading": "Daily attendance (P = present, blank = absent)",
        "no": "No",
        "national_id": "National ID",
        "name": "Name",
        "grade": "Class",
        "parent_phone": "Parent Phone",
        "days_present": "Days Present",
        "days_absent": "Days Absent",
        "pct": "Attendance %",
        "present": "Present",
        "absent": "Absent",
        "total": "Total",
        "date": "Date",
        "time": "Time",
        "method": "Method",
        "method_face": "face",
        "method_manual": "manual",
        "mark_present": "P",
        "mark_absent": "A",
    },
}


def current_lang():
    lang = request.args.get("lang") or request.headers.get("X-Lang") or "ar"
    return lang if lang in STRINGS else "ar"


def msg(key, **kw):
    return STRINGS[current_lang()][key].format(**kw)


_ARABIC_RE = re.compile(r"[\u0600-\u06FF]")


def shape(text):
    """Shape Arabic text for PDF rendering (reportlab does not do bidi/joining itself)."""
    text = "" if text is None else str(text)
    if not _ARABIC_RE.search(text):
        return text
    return get_display(arabic_reshaper.reshape(text))


_pdf_fonts_ready = False


def register_pdf_fonts():
    global _pdf_fonts_ready
    if _pdf_fonts_ready:
        return
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.pdfmetrics import registerFontFamily
    from reportlab.pdfbase.ttfonts import TTFont

    pdfmetrics.registerFont(TTFont("Amiri", os.path.join(FONT_DIR, "Amiri-Regular.ttf")))
    pdfmetrics.registerFont(TTFont("Amiri-Bold", os.path.join(FONT_DIR, "Amiri-Bold.ttf")))
    registerFontFamily("Amiri", normal="Amiri", bold="Amiri-Bold", italic="Amiri", boldItalic="Amiri-Bold")
    _pdf_fonts_ready = True


# --------------------------------------------------------------------------- db
def get_db():
    if "db" not in g:
        conn = sqlite3.connect(DB_PATH)
        conn.row_factory = sqlite3.Row
        conn.execute("PRAGMA journal_mode=WAL")
        conn.execute("PRAGMA foreign_keys=ON")
        g.db = conn
    return g.db


@app.teardown_appcontext
def close_db(_exc):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def init_db():
    conn = sqlite3.connect(DB_PATH)
    conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS students (
            id          INTEGER PRIMARY KEY AUTOINCREMENT,
            national_id TEXT NOT NULL UNIQUE,
            name        TEXT NOT NULL,
            grade       TEXT NOT NULL DEFAULT '',
            bus_no      TEXT NOT NULL DEFAULT '',
            parent_phone TEXT NOT NULL DEFAULT '',
            photo       TEXT NOT NULL DEFAULT '',
            descriptors TEXT NOT NULL DEFAULT '[]',
            created_at  TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS attendance (
            id         INTEGER PRIMARY KEY AUTOINCREMENT,
            student_id INTEGER NOT NULL REFERENCES students(id) ON DELETE CASCADE,
            day        TEXT NOT NULL,
            time       TEXT NOT NULL,
            method     TEXT NOT NULL DEFAULT 'face',
            UNIQUE(student_id, day)
        );
        CREATE INDEX IF NOT EXISTS idx_attendance_day ON attendance(day);
        """
    )
    cols = {r[1] for r in conn.execute("PRAGMA table_info(students)")}
    if "student_no" in cols and "national_id" not in cols:
        conn.execute("ALTER TABLE students RENAME COLUMN student_no TO national_id")
    if "bus_no" not in cols:
        conn.execute("ALTER TABLE students ADD COLUMN bus_no TEXT NOT NULL DEFAULT ''")
    if "parent_phone" not in cols:
        conn.execute("ALTER TABLE students ADD COLUMN parent_phone TEXT NOT NULL DEFAULT ''")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_students_bus ON students(bus_no)")
    conn.commit()
    conn.close()


init_db()


# ---------------------------------------------------------------------- helpers
def today_str():
    return date.today().isoformat()


def valid_day(value, default=None):
    if not value:
        return default
    if not DAY_RE.match(value):
        return None
    try:
        date.fromisoformat(value)
    except ValueError:
        return None
    return value


def student_row_to_dict(row, include_descriptors=True):
    data = {
        "id": row["id"],
        "national_id": row["national_id"],
        "name": row["name"],
        "grade": row["grade"],
        "bus_no": row["bus_no"],
        "parent_phone": row["parent_phone"],
        "photo": row["photo"],
        "created_at": row["created_at"],
    }
    if include_descriptors:
        data["descriptors"] = json.loads(row["descriptors"] or "[]")
    return data


def clean_descriptors(raw):
    """Accept a list of 128-float vectors; return a JSON string or raise ValueError."""
    if raw is None:
        return "[]"
    if not isinstance(raw, list):
        raise ValueError(msg("desc_list"))
    out = []
    for vec in raw:
        if not isinstance(vec, list) or len(vec) != 128:
            raise ValueError(msg("desc_len"))
        out.append([float(x) for x in vec])
    if len(out) > 10:
        out = out[:10]
    return json.dumps(out)


def range_from_args():
    """Read from/to query params. Defaults to the current month."""
    today = date.today()
    default_from = today.replace(day=1).isoformat()
    default_to = today.isoformat()
    start = valid_day(request.args.get("from"), default_from)
    end = valid_day(request.args.get("to"), default_to)
    if start is None or end is None:
        return None, None
    if start > end:
        start, end = end, start
    return start, end


# ---------------------------------------------------------------------- routes
@app.route("/")
def index():
    return send_from_directory(app.static_folder, "index.html")


@app.route("/health")
def health():
    return jsonify({"ok": True})


@app.route("/api/students", methods=["GET"])
def list_students():
    db = get_db()
    day = valid_day(request.args.get("date"), today_str())
    rows = db.execute("SELECT * FROM students ORDER BY name COLLATE NOCASE").fetchall()
    present = {
        r["student_id"]: r["time"]
        for r in db.execute("SELECT student_id, time FROM attendance WHERE day = ?", (day,))
    }
    counts = {
        r["student_id"]: r["n"]
        for r in db.execute("SELECT student_id, COUNT(*) AS n FROM attendance GROUP BY student_id")
    }
    students = []
    for row in rows:
        s = student_row_to_dict(row)
        s["present_today"] = row["id"] in present
        s["time_today"] = present.get(row["id"])
        s["total_days"] = counts.get(row["id"], 0)
        students.append(s)
    return jsonify({"date": day, "students": students})


@app.route("/api/buses", methods=["GET"])
def list_buses():
    db = get_db()
    rows = db.execute(
        "SELECT bus_no, COUNT(*) AS n FROM students WHERE bus_no <> '' GROUP BY bus_no ORDER BY bus_no COLLATE NOCASE"
    ).fetchall()
    return jsonify({"buses": [{"bus_no": r["bus_no"], "students": r["n"]} for r in rows]})


@app.route("/api/students", methods=["POST"])
def create_student():
    payload = request.get_json(silent=True) or {}
    name = (payload.get("name") or "").strip()
    national_id = (payload.get("national_id") or "").strip()
    grade = (payload.get("grade") or "").strip()
    bus_no = (payload.get("bus_no") or "").strip()
    parent_phone = (payload.get("parent_phone") or "").strip()
    photo = payload.get("photo") or ""
    if not name:
        return jsonify({"error": msg("name_required")}), 400
    if not national_id:
        return jsonify({"error": msg("id_required")}), 400
    if photo and not photo.startswith("data:image/"):
        return jsonify({"error": msg("photo_image")}), 400
    if len(photo) > 400_000:
        return jsonify({"error": msg("photo_large")}), 400
    try:
        descriptors = clean_descriptors(payload.get("descriptors"))
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400

    db = get_db()
    try:
        cur = db.execute(
            "INSERT INTO students (national_id, name, grade, bus_no, parent_phone, photo, descriptors, created_at) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
            (national_id, name, grade, bus_no, parent_phone, photo, descriptors, datetime.now().isoformat(timespec="seconds")),
        )
        db.commit()
    except sqlite3.IntegrityError:
        return jsonify({"error": msg("id_exists", id=national_id)}), 409
    row = db.execute("SELECT * FROM students WHERE id = ?", (cur.lastrowid,)).fetchone()
    s = student_row_to_dict(row)
    s.update({"present_today": False, "time_today": None, "total_days": 0})
    return jsonify(s), 201


@app.route("/api/students/<int:student_id>", methods=["PUT"])
def update_student(student_id):
    payload = request.get_json(silent=True) or {}
    db = get_db()
    row = db.execute("SELECT * FROM students WHERE id = ?", (student_id,)).fetchone()
    if row is None:
        return jsonify({"error": msg("not_found")}), 404

    name = (payload.get("name") or row["name"]).strip()
    national_id = (payload.get("national_id") or row["national_id"]).strip()
    grade = (payload.get("grade") if payload.get("grade") is not None else row["grade"]).strip()
    bus_no = (payload.get("bus_no") if payload.get("bus_no") is not None else row["bus_no"]).strip()
    parent_phone = (payload.get("parent_phone") if payload.get("parent_phone") is not None else row["parent_phone"]).strip()
    photo = payload.get("photo") if payload.get("photo") is not None else row["photo"]
    if photo and not photo.startswith("data:image/"):
        return jsonify({"error": msg("photo_image")}), 400
    descriptors = row["descriptors"]
    if "descriptors" in payload:
        try:
            descriptors = clean_descriptors(payload.get("descriptors"))
        except ValueError as exc:
            return jsonify({"error": str(exc)}), 400
    try:
        db.execute(
            "UPDATE students SET national_id=?, name=?, grade=?, bus_no=?, parent_phone=?, photo=?, descriptors=? WHERE id=?",
            (national_id, name, grade, bus_no, parent_phone, photo, descriptors, student_id),
        )
        db.commit()
    except sqlite3.IntegrityError:
        return jsonify({"error": msg("id_exists", id=national_id)}), 409
    row = db.execute("SELECT * FROM students WHERE id = ?", (student_id,)).fetchone()
    return jsonify(student_row_to_dict(row))


@app.route("/api/students/<int:student_id>", methods=["DELETE"])
def delete_student(student_id):
    db = get_db()
    cur = db.execute("DELETE FROM students WHERE id = ?", (student_id,))
    db.commit()
    if cur.rowcount == 0:
        return jsonify({"error": msg("not_found")}), 404
    return jsonify({"ok": True})


@app.route("/api/students/<int:student_id>/history", methods=["GET"])
def student_history(student_id):
    db = get_db()
    row = db.execute("SELECT * FROM students WHERE id = ?", (student_id,)).fetchone()
    if row is None:
        return jsonify({"error": msg("not_found")}), 404
    records = db.execute(
        "SELECT id, day, time, method FROM attendance WHERE student_id = ? ORDER BY day DESC LIMIT 90",
        (student_id,),
    ).fetchall()
    return jsonify(
        {
            "student": student_row_to_dict(row, include_descriptors=False),
            "records": [dict(r) for r in records],
        }
    )


@app.route("/api/attendance", methods=["GET"])
def list_attendance():
    db = get_db()
    day = valid_day(request.args.get("date"), today_str())
    if day is None:
        return jsonify({"error": msg("invalid_date")}), 400
    rows = db.execute(
        "SELECT a.id, a.day, a.time, a.method, s.id AS student_id, s.national_id, s.name, s.grade, s.bus_no, s.parent_phone, s.photo "
        "FROM attendance a JOIN students s ON s.id = a.student_id "
        "WHERE a.day = ? ORDER BY a.time DESC",
        (day,),
    ).fetchall()
    total = db.execute("SELECT COUNT(*) FROM students").fetchone()[0]
    return jsonify({"date": day, "total_students": total, "records": [dict(r) for r in rows]})


@app.route("/api/attendance", methods=["POST"])
def mark_attendance():
    payload = request.get_json(silent=True) or {}
    student_id = payload.get("student_id")
    if not isinstance(student_id, int):
        return jsonify({"error": msg("student_id_required")}), 400
    day = valid_day(payload.get("day"), today_str())
    if day is None:
        return jsonify({"error": msg("invalid_date")}), 400
    time_str = payload.get("time") or datetime.now().strftime("%H:%M:%S")
    if not re.match(r"^\d{2}:\d{2}(:\d{2})?$", str(time_str)):
        time_str = datetime.now().strftime("%H:%M:%S")
    method = payload.get("method") if payload.get("method") in ("face", "manual") else "face"

    db = get_db()
    student = db.execute("SELECT * FROM students WHERE id = ?", (student_id,)).fetchone()
    if student is None:
        return jsonify({"error": msg("not_found")}), 404
    existing = db.execute(
        "SELECT id, time FROM attendance WHERE student_id = ? AND day = ?", (student_id, day)
    ).fetchone()
    if existing:
        return jsonify(
            {
                "already_marked": True,
                "id": existing["id"],
                "student_id": student_id,
                "name": student["name"],
                "national_id": student["national_id"],
                "day": day,
                "time": existing["time"],
            }
        )
    cur = db.execute(
        "INSERT INTO attendance (student_id, day, time, method) VALUES (?, ?, ?, ?)",
        (student_id, day, time_str, method),
    )
    db.commit()
    return (
        jsonify(
            {
                "already_marked": False,
                "id": cur.lastrowid,
                "student_id": student_id,
                "name": student["name"],
                "national_id": student["national_id"],
                "day": day,
                "time": time_str,
            }
        ),
        201,
    )


@app.route("/api/attendance/<int:record_id>", methods=["DELETE"])
def delete_attendance(record_id):
    db = get_db()
    cur = db.execute("DELETE FROM attendance WHERE id = ?", (record_id,))
    db.commit()
    if cur.rowcount == 0:
        return jsonify({"error": msg("record_not_found")}), 404
    return jsonify({"ok": True})


# ---------------------------------------------------------------------- export
def build_report(start, end, bus_no=""):
    """Collect everything the exporters need for the given inclusive date range (optionally one bus)."""
    db = get_db()
    if bus_no:
        students = db.execute(
            "SELECT id, national_id, name, grade, bus_no, parent_phone FROM students WHERE bus_no = ? ORDER BY name COLLATE NOCASE",
            (bus_no,),
        ).fetchall()
        records = db.execute(
            "SELECT a.student_id, a.day, a.time, a.method FROM attendance a JOIN students s ON s.id = a.student_id "
            "WHERE s.bus_no = ? AND a.day BETWEEN ? AND ? ORDER BY a.day, a.time",
            (bus_no, start, end),
        ).fetchall()
    else:
        students = db.execute(
            "SELECT id, national_id, name, grade, bus_no, parent_phone FROM students ORDER BY name COLLATE NOCASE"
        ).fetchall()
        records = db.execute(
            "SELECT a.student_id, a.day, a.time, a.method FROM attendance a "
            "WHERE a.day BETWEEN ? AND ? ORDER BY a.day, a.time",
            (start, end),
        ).fetchall()
    days = sorted({r["day"] for r in records})
    by_student = {}
    for r in records:
        by_student.setdefault(r["student_id"], {})[r["day"]] = r["time"]
    return {
        "start": start,
        "end": end,
        "bus_no": bus_no,
        "lang": current_lang(),
        "T": STRINGS[current_lang()],
        "students": [dict(s) for s in students],
        "records": [dict(r) for r in records],
        "days": days,
        "by_student": by_student,
        "generated": datetime.now().strftime("%Y-%m-%d %H:%M"),
    }


def export_filename(report, ext):
    bus = f"_bus-{re.sub(r'[^A-Za-z0-9-]+', '', report['bus_no'])}" if report["bus_no"] else ""
    return f"bus-attendance{bus}_{report['start']}_to_{report['end']}.{ext}"


def bus_from_args():
    return (request.args.get("bus") or "").strip()[:40]


@app.route("/api/export/excel")
def export_excel():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    start, end = range_from_args()
    if start is None:
        return jsonify({"error": msg("invalid_range")}), 400
    report = build_report(start, end, bus_from_args())
    days = report["days"]
    n_days = len(days)
    T = report["T"]
    rtl = report["lang"] == "ar"
    bus_title = f"  {T['bus']} {report['bus_no']}" if report["bus_no"] else ""

    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F4E78")
    present_fill = PatternFill("solid", fgColor="C6EFCE")
    absent_fill = PatternFill("solid", fgColor="FFC7CE")
    center = Alignment(horizontal="center")

    def style_header(ws, row=1):
        for cell in ws[row]:
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = center

    # Sheet 1: Summary
    ws = wb.active
    ws.title = T["sheet_summary"]
    ws.sheet_view.rightToLeft = rtl
    ws.append([f"{T['summary_title']}{bus_title}  ({start} - {end})"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"{T['school_days']}: {n_days}    {T['students']}: {len(report['students'])}    {T['generated']}: {report['generated']}"])
    ws.append([])
    ws.append([T["no"], T["national_id"], T["name"], T["grade"], T["bus"], T["parent_phone"], T["days_present"], T["days_absent"], T["pct"]])
    style_header(ws, 4)
    for i, s in enumerate(report["students"], 1):
        present = len(report["by_student"].get(s["id"], {}))
        pct = round(100 * present / n_days, 1) if n_days else 0
        ws.append([i, s["national_id"], s["name"], s["grade"], s["bus_no"], s["parent_phone"], present, n_days - present, pct])
    for col, width in zip("ABCDEFGHI", (6, 18, 30, 12, 10, 18, 14, 14, 14)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A5"

    # Sheet 2: Daily matrix
    ws2 = wb.create_sheet(T["sheet_daily"])
    ws2.sheet_view.rightToLeft = rtl
    header = [T["national_id"], T["name"], T["grade"], T["bus"]] + days + [T["total"]]
    ws2.append(header)
    style_header(ws2)
    for s in report["students"]:
        marks = report["by_student"].get(s["id"], {})
        row = [s["national_id"], s["name"], s["grade"], s["bus_no"]] + [T["mark_present"] if d in marks else T["mark_absent"] for d in days] + [len(marks)]
        ws2.append(row)
        r = ws2.max_row
        for j, d in enumerate(days, start=5):
            cell = ws2.cell(row=r, column=j)
            cell.alignment = center
            cell.fill = present_fill if d in marks else absent_fill
    ws2.column_dimensions["A"].width = 18
    ws2.column_dimensions["B"].width = 30
    ws2.column_dimensions["C"].width = 12
    ws2.column_dimensions["D"].width = 10
    for j in range(5, 5 + n_days):
        ws2.column_dimensions[get_column_letter(j)].width = 11
    ws2.freeze_panes = "E2"

    # Sheet 3: Log
    ws3 = wb.create_sheet(T["sheet_log"])
    ws3.sheet_view.rightToLeft = rtl
    ws3.append([T["date"], T["time"], T["national_id"], T["name"], T["grade"], T["bus"], T["method"]])
    style_header(ws3)
    lookup = {s["id"]: s for s in report["students"]}
    for r in report["records"]:
        s = lookup.get(r["student_id"], {})
        ws3.append([r["day"], r["time"], s.get("national_id", ""), s.get("name", ""), s.get("grade", ""), s.get("bus_no", ""), T.get("method_" + r["method"], r["method"])])
    for col, width in zip("ABCDEFG", (12, 10, 18, 30, 12, 10, 10)):
        ws3.column_dimensions[col].width = width
    ws3.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name=export_filename(report, "xlsx"),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/api/export/pdf")
def export_pdf():
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    start, end = range_from_args()
    if start is None:
        return jsonify({"error": msg("invalid_range")}), 400
    report = build_report(start, end, bus_from_args())
    days = report["days"]
    n_days = len(days)
    T = report["T"]
    rtl = report["lang"] == "ar"
    bus_title = f" – {T['bus']} {report['bus_no']}" if report["bus_no"] else ""
    register_pdf_fonts()
    font, font_bold = "Amiri", "Amiri-Bold"

    def cells(row):
        """Shape every cell and, for Arabic, reverse the column order so the first column sits on the right."""
        out = [shape(c) for c in row]
        return out[::-1] if rtl else out

    def widths(ws):
        return ws[::-1] if rtl else ws

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=T["report_title"],
    )
    styles = getSampleStyleSheet()
    for name in ("Title", "Normal", "Heading2"):
        styles[name].fontName = font_bold if name != "Normal" else font
        if rtl and name != "Title":
            styles[name].alignment = 2  # right
    meta = (
        f"{T['period']}: {start} - {end}   |   {T['school_days']}: {n_days}   |   "
        f"{T['students']}: {len(report['students'])}   |   {T['generated']}: {report['generated']}"
    )
    story = [
        Paragraph(shape(f"{T['report_title']}{bus_title}"), styles["Title"]),
        Paragraph(shape(meta), styles["Normal"]),
        Spacer(1, 6 * mm),
    ]

    header_style = TableStyle(
        [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, -1), font),
            ("FONTNAME", (0, 0), (-1, 0), font_bold),
            ("ALIGN", (0, 0), (-1, -1), "RIGHT" if rtl else "LEFT"),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F2")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]
    )

    # Summary table
    data = [cells([T["no"], T["national_id"], T["name"], T["grade"], T["bus"], T["parent_phone"], T["present"], T["absent"], "%"])]
    for i, s in enumerate(report["students"], 1):
        present = len(report["by_student"].get(s["id"], {}))
        pct = f"{100 * present / n_days:.0f}%" if n_days else "-"
        data.append(cells([i, s["national_id"], s["name"], s["grade"], s["bus_no"], s["parent_phone"], present, n_days - present, pct]))
    table = Table(data, repeatRows=1, colWidths=widths([10 * mm, 34 * mm, 70 * mm, 20 * mm, 16 * mm, 34 * mm, 20 * mm, 20 * mm, 16 * mm]))
    table.setStyle(header_style)
    table.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 8)]))
    story.append(Paragraph(shape(T["summary_heading"]), styles["Heading2"]))
    story.append(table)

    # Daily matrix (fits comfortably up to a month of school days)
    if 0 < n_days <= 31:
        story.append(PageBreak())
        story.append(Paragraph(shape(T["daily_heading"]), styles["Heading2"]))
        day_labels = [d[5:] for d in days]  # MM-DD
        mdata = [cells([T["national_id"], T["name"], T["bus"]] + day_labels + [T["total"]])]
        for s in report["students"]:
            marks = report["by_student"].get(s["id"], {})
            mdata.append(cells([s["national_id"], s["name"], s["bus_no"]] + [T["mark_present"] if d in marks else "" for d in days] + [len(marks)]))
        avail = landscape(A4)[0] - 24 * mm - 32 * mm - 50 * mm - 14 * mm - 10 * mm
        day_w = min(12 * mm, avail / n_days)
        mtable = Table(mdata, repeatRows=1, colWidths=widths([32 * mm, 50 * mm, 14 * mm] + [day_w] * n_days + [10 * mm]))
        mtable.setStyle(header_style)
        # the day columns sit after the 3 identity columns (or before the mirrored ones in RTL)
        day_first, day_last = (1, n_days) if rtl else (3, 2 + n_days)
        mtable.setStyle(
            TableStyle(
                [
                    ("FONTSIZE", (0, 0), (-1, -1), 6),
                    ("ALIGN", (day_first, 0), (day_last, -1), "CENTER"),
                    ("TEXTCOLOR", (day_first, 1), (day_last, -1), colors.HexColor("#0B6623")),
                ]
            )
        )
        story.append(mtable)

    doc.build(story)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=export_filename(report, "pdf"), mimetype="application/pdf")


@app.after_request
def add_headers(resp):
    if request.path.startswith("/static/models/") or request.path.startswith("/static/vendor/"):
        resp.headers["Cache-Control"] = "public, max-age=31536000, immutable"
    elif request.path.startswith("/api/"):
        resp.headers["Cache-Control"] = "no-store"
    return resp


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 8080)), debug=True)
